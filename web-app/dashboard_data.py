"""Dashboard de inicio: arma el JSON y el Excel para persona y empresa.

Los dos portales muestran exactamente lo mismo; sólo cambia el scope
(`{"individual_id": …}` o `{"business_id": …}`), así que la lógica vive acá
una sola vez y las rutas de cada blueprint quedan de dos líneas.

Todos los montos van en colones (amount_local) y salen de tools/finance.py,
que ya excluye descartados, duplicados y no aprobados.
"""

import calendar
import io
from datetime import date, timedelta

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tools import finance

MESES_HISTORICO = 12   # el histórico mira hasta 12 meses atrás…
MESES_DETALLE = 6      # …y el detalle de una categoría, 6 (mes actual incluido)


def _primer_dia_meses_atras(today, n):
    """Primer día del mes que está n-1 meses antes del mes de `today`
    (n=1 → este mes; n=12 → doce meses contando el actual)."""
    idx = today.year * 12 + (today.month - 1) - (n - 1)
    return date(idx // 12, idx % 12 + 1, 1)


def _historico_desde(conn, scope, today):
    """Arranque del histórico: 12 meses atrás, pero nunca antes del mes de la
    primera transacción del cliente — para no dibujar meses vacíos."""
    desde = _primer_dia_meses_atras(today, MESES_HISTORICO)
    primera = finance.get_first_transaction_date(conn, **scope)
    if primera and primera > desde:
        desde = primera.replace(day=1)
    return desde, primera


def resumen(conn, scope, today=None):
    """Payload de la carga inicial: mes en curso, composición, presupuestos,
    histórico. Una sola llamada para que la página aparezca completa."""
    today = today or date.today()
    mes_desde = today.replace(day=1)
    dias_mes = calendar.monthrange(today.year, today.month)[1]
    hist_desde, primera = _historico_desde(conn, scope, today)

    # Mismo tramo (día 1 al día de hoy) del mes anterior, para el delta.
    prev_fin = mes_desde - timedelta(days=1)
    prev_ini = prev_fin.replace(day=1)
    prev_hasta = prev_ini.replace(day=min(today.day, prev_fin.day))

    return {
        "hoy": str(today),
        "mes": {
            "desde": str(mes_desde),
            "hasta": str(today),
            "dias_restantes": dias_mes - today.day,
            "actual": finance.get_income_expense_summary(
                conn, **scope, date_from=mes_desde, date_to=today),
            "mismo_tramo_anterior": finance.get_income_expense_summary(
                conn, **scope, date_from=prev_ini, date_to=prev_hasta),
        },
        "categorias": finance.get_top_spending(
            conn, **scope, date_from=mes_desde, date_to=today, limit=10_000),
        "presupuestos": finance.get_budget_status(
            conn, **scope, date_from=mes_desde, date_to=today),
        "historico": {
            "desde": str(hist_desde),
            "primera_transaccion": str(primera) if primera else None,
            "serie": finance.get_monthly_income_expense(
                conn, **scope, date_from=hist_desde, date_to=today),
        },
    }


def categoria(conn, scope, category, subcategory, today=None):
    """Detalle de una Categoría/Subcategoría: últimos 6 meses, presupuesto y
    los comercios donde más se gastó en lo que va del mes."""
    today = today or date.today()
    mes_desde = today.replace(day=1)
    return {
        "category": category,
        "subcategory": subcategory,
        "serie": finance.get_monthly_category_spending(
            conn, **scope, category=category, subcategory=subcategory,
            date_from=_primer_dia_meses_atras(today, MESES_DETALLE), date_to=today),
        "budget": finance.get_category_budget(
            conn, **scope, category=category, subcategory=subcategory),
        "comercios": finance.get_top_merchants(
            conn, **scope, category=category, subcategory=subcategory,
            date_from=mes_desde, date_to=today, limit=3),
    }


# ── Excel ────────────────────────────────────────────────────────────────────

_INK = "1B1C20"
_FILL_HEAD = PatternFill("solid", fgColor=_INK)
_FONT_HEAD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
_FONT_BOLD = Font(name="Calibri", size=11, bold=True)
_CRC_FMT = '#,##0.00'
_PCT_FMT = '0.0%'
_TIPO_LABEL = {"credito": "Ingreso", "debito": "Gasto"}


def _encabezado(ws, headers):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = _FONT_HEAD
        cell.fill = _FILL_HEAD
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _ajustar_anchos(ws, minimo=10, maximo=40):
    for i, col in enumerate(ws.columns, 1):
        largo = max((len(str(c.value)) for c in col if c.value is not None), default=minimo)
        ws.column_dimensions[get_column_letter(i)].width = min(max(largo + 2, minimo), maximo)


def excel(conn, scope, today=None):
    """Libro con dos hojas: Resumen (mes a mes) y Por categoría (una fila por
    categoría/subcategoría y tipo, una columna por mes). Devuelve
    (BytesIO, nombre_de_archivo). Mismo período que el histórico en pantalla."""
    today = today or date.today()
    desde, _ = _historico_desde(conn, scope, today)
    meses = finance.month_buckets(desde, today)

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Resumen"
    _encabezado(ws, ["Mes", "Ingresos", "Gastos", "Ahorro", "Tasa de ahorro"])
    for h in finance.get_monthly_income_expense(conn, **scope, date_from=desde, date_to=today):
        ws.append([h["month"], h["ingresos"], h["gastos"], h["ingresos"] - h["gastos"],
                   None if h["tasa_ahorro"] is None else h["tasa_ahorro"] / 100])
    for fila in ws.iter_rows(min_row=2):
        for cell in fila[1:4]:
            cell.number_format = _CRC_FMT
        fila[4].number_format = _PCT_FMT
    _ajustar_anchos(ws)

    ws2 = wb.create_sheet("Por categoría")
    _encabezado(ws2, ["Tipo", "Categoría", "Subcategoría", *meses, "Total"])
    for r in finance.get_category_month_matrix(conn, **scope, date_from=desde, date_to=today):
        valores = [r["by_month"].get(m, 0.0) for m in meses]
        ws2.append([_TIPO_LABEL.get(r["tipo"], r["tipo"]), r["category"],
                    r["subcategory"] or "", *valores, sum(valores)])
    for fila in ws2.iter_rows(min_row=2):
        for cell in fila[3:]:
            cell.number_format = _CRC_FMT
        fila[-1].font = _FONT_BOLD
    _ajustar_anchos(ws2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, f"dashboard_{desde}_{today}.xlsx"
